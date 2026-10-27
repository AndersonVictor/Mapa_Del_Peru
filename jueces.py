"""Datos del dashboard de la CSJ Junín, desde dashboard.xlsx (tablas planas, una fila por registro).

- Magistrados -> un juez o jueza por fila: tarjetas, gráficos del resumen y burbujas por sede.
- Organos     -> un órgano jurisdiccional por fila (tarjeta de órganos).
- El resto de hojas (Poblacion, Carga_Procesal, Audiencias_*, Personal...) se envían tal cual y
  el navegador hace las sumas, así los totales salen siempre de la misma fuente.

Los nombres de provincia y distrito se cruzan con datos.xlsx sin importar tildes ni Ñ.
"""
import datetime
import re
import unicodedata
from collections import defaultdict

CAMPOS = ["total", "varones", "mujeres", "titulares", "prov", "supern", "sala", "jesp", "jpl"]


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).upper()
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def limpio(v):
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.isoformat()
    return v


def leer_hojas(ruta):
    """{hoja: [ {COLUMNA: valor}, ... ]}. Encabezado en la fila 1; se saltan filas vacías."""
    from openpyxl import load_workbook
    wb = load_workbook(ruta, read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        filas = ws.iter_rows(values_only=True)
        cab = [str(c).strip() if c is not None else "" for c in next(filas, [])]
        out[ws.title] = [{c: limpio(v) for c, v in zip(cab, f) if c}
                         for f in filas if any(v not in (None, "") for v in f)]
    return out


class Buscador:
    """Nombre (+ departamento/provincia) -> id_geografia de datos.xlsx. Estricto: si no calza, None."""

    def __init__(self, datos):
        self.por_nivel = defaultdict(list)
        for gid, d in datos.items():
            campo = {"7": "nombregion", "8": "nombprov", "9": "nombdist"}.get(gid[0])
            if campo:
                self.por_nivel[gid[0]].append((gid, norm(d.get(campo)), norm(d.get("nombregion")), norm(d.get("nombprov"))))

    def __call__(self, nivel, nombre, dep=None, prov=None):
        n, dep, prov = norm(nombre), norm(dep), norm(prov)
        cand = [c for c in self.por_nivel[nivel] if (not dep or c[2] == dep) and (not prov or c[3] == prov)]
        exacto = [c[0] for c in cand if c[1] == n]
        if exacto:
            return exacto[0]
        # 'SAN AGUSTIN DE CAJAS' vs 'SAN AGUSTIN': vale si es el único dentro de la misma provincia
        parecido = [c[0] for c in cand if c[1] and (n.startswith(c[1] + " ") or c[1].startswith(n + " "))]
        return parecido[0] if len(parecido) == 1 else None


def sumar(destino, m):
    sexo, cond, inst = norm(m.get("SEXO")), norm(m.get("CONDICION")), norm(m.get("INSTANCIA"))
    destino["total"] += 1
    destino["varones"] += sexo.startswith("VAR")
    destino["mujeres"] += sexo.startswith("MUJ")
    destino["titulares"] += cond.startswith("T")
    destino["prov"] += cond.startswith("P")
    destino["supern"] += cond.startswith("S")
    destino["sala"] += inst == "SALA"
    destino["jesp"] += inst.startswith("J ESP")
    destino["jpl"] += inst == "JPL"


def resumen(ruta, datos, cajas):
    """{'tablas': {hoja: filas}, 'provincias': {id: {...}}, 'sedes': [...], 'csjj': {...}, 'avisos': [...]}"""
    tablas = leer_hojas(ruta)
    buscar = Buscador(datos)
    avisos = []
    vacio = lambda **k: dict({c: 0 for c in CAMPOS}, **k)

    # Cada magistrado: id de su provincia y de su distrito (sede) en el mapa
    provincias, sedes, sin_ubicar = {}, {}, defaultdict(int)
    for m in tablas.get("Magistrados", []):
        dep, prov, sede = m.get("DEPARTAMENTO"), m.get("PROVINCIA"), m.get("SEDE_DISTRITO")
        pid = buscar("8", prov, dep)
        did = buscar("9", sede, dep, prov)
        m["prov_id"], m["dist_id"] = pid, did
        if pid:
            sumar(provincias.setdefault(pid, vacio(nombre=prov, departamento=dep or "")), m)
        else:
            avisos.append(f"Magistrados: provincia '{prov}' ({dep}) no se encontró en datos.xlsx.")
        if did:
            sumar(sedes.setdefault(did, vacio(nombre=sede, provincia=prov, departamento=dep or "", id=did)), m)
        else:
            sin_ubicar[f"{sede} ({prov})"] += 1
    csjj = vacio()
    for m in tablas.get("Magistrados", []):
        sumar(csjj, m)

    # Órganos: provincia y distrito en el mapa (para filtrar la tarjeta de órganos)
    for o in tablas.get("Organos", []):
        o["prov_id"] = buscar("8", o.get("PROVINCIA"), o.get("DEPARTAMENTO"))
        o["dist_id"] = buscar("9", o.get("DISTRITO"), o.get("DEPARTAMENTO"), o.get("PROVINCIA"))

    # Burbujas: en el centro de su distrito (el mapa ya sabe dónde está cada distrito)
    for did, s in sedes.items():
        c = cajas.get(did)
        if c:
            s["lon"], s["lat"] = (c[4], c[5]) if len(c) >= 6 else ((c[0] + c[2]) / 2, (c[1] + c[3]) / 2)

    for k, n in sin_ubicar.items():
        avisos.append(f"Magistrados: sede '{k}' no coincide con un distrito de datos.xlsx ({n} juez/jueza sin burbuja).")

    return {"tablas": tablas, "provincias": provincias, "sedes": list(sedes.values()), "csjj": csjj, "avisos": avisos}
