"""Territorio de la Corte leído de competencias.xlsx (columna 'Competencia (Territorial)').

Cada texto se interpreta así:
  - "PROVINCIA(S) DE A, B Y C"            -> provincias completas
  - "DISTRITO(S) DE X, Y Y Z"             -> solo esos distritos
  - "(PROVINCIA DE ...)", "DEPARTAMENTO DE ...", "EXCEPCIÓN ...", "DISTRITO JUDICIAL ..." -> se ignoran
  - Alias: LA OROYA = provincia Yauli, PAMPAS = provincia Tayacaja, CAJAS = San Agustín (de Cajas)
  - Errores de tipeo (CHICHIHUASI, PUNCO, YROBLE...) se corrigen con el nombre más parecido.
Un distrito se busca primero en la provincia de la sede del órgano, luego en las provincias del mismo
texto y por último en Junín / Huancavelica.
"""
import difflib
import re
import unicodedata
from collections import defaultdict

DEPARTAMENTOS = ("JUNIN", "HUANCAVELICA")
ALIAS_PROVINCIA = {"LA OROYA": "YAULI", "PAMPAS": "TAYACAJA"}
ALIAS_DISTRITO = {"CAJAS": "SAN AGUSTIN"}


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).upper()
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


class Geo:
    """Provincias y distritos de Junín y Huancavelica (de datos.xlsx)."""

    def __init__(self, datos):
        self.prov = {}                      # nombre -> id
        self.dist = defaultdict(list)       # nombre -> [(id, id_provincia)]
        self.nombre = {}
        for gid, d in datos.items():
            if norm(d.get("nombregion")) not in DEPARTAMENTOS:
                continue
            if gid[0] == "8":
                self.prov[norm(d.get("nombprov"))] = gid
                self.nombre[gid] = d.get("nombprov")
            elif gid[0] == "9":
                self.dist[norm(d.get("nombdist"))].append((gid, "8" + gid[1:5] + "00"))
                self.nombre[gid] = d.get("nombdist")

    def parecido(self, nombre, opciones):
        r = difflib.get_close_matches(nombre, list(opciones), n=1, cutoff=0.82)
        return r[0] if r else None

    def provincia(self, nombre):
        nombre = ALIAS_PROVINCIA.get(nombre, nombre)
        if nombre in self.prov:
            return self.prov[nombre]
        p = self.parecido(nombre, self.prov)
        return self.prov[p] if p else None

    def distrito(self, nombre, preferidas):
        nombre = ALIAS_DISTRITO.get(nombre, nombre)
        cand = self.dist.get(nombre) or self.dist.get(self.parecido(nombre, self.dist) or "", [])
        if not cand:
            return None
        for p in preferidas:                # primero en la provincia de la sede o del mismo texto
            for gid, prov in cand:
                if prov == p:
                    return gid
        return cand[0][0] if len(cand) == 1 else None


def partes(texto):
    """Texto -> trozos con su modo ('prov' o 'dist')."""
    # mayúsculas y sin tildes, pero conservando comas, puntos y guiones (separan los nombres)
    t = unicodedata.normalize("NFD", str(texto)).upper()
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    t = re.sub(r"\s+", " ", t)
    # "(PROVINCIA DE CHURCAMPA)" y "DE LA PROVINCIA DE TAYACAJA" solo aclaran dónde están los distritos
    for m in re.finditer(r"(?:\(\s*|DE LA )PROVINCIA DE ([A-Z ]+)", t):
        yield "ctx", norm(m.group(1))
    t = re.sub(r"\([^)]*\)", " ", t)
    t = re.sub(r"\bDE LA PROVINCIA DE [A-Z ]+", " ", t)
    t = re.sub(r"(EN TODO EL )?\bDISTRITOS? JUDICIAL(ES)?\b.*", " ", t)   # distritos judiciales: no es geografía
    t = re.sub(r"[^A-Z0-9 ,.;-]", " ", t)
    t = re.sub(r"\bEXCEPCION\b.*", " ", t)
    t = re.sub(r"\bDEPARTAMENTO DE [A-Z]+", " ", t)
    modo = "prov"
    for trozo in re.split(r",| Y |\.| ?- ?|;", " " + t + " "):
        trozo = norm(trozo)
        if not trozo or "JUDICIAL" in trozo:
            continue
        while True:                                            # quita y recuerda palabras guía
            m = re.match(r"^(Y |LOS |LAS |EN TODO EL |DE )", trozo)
            if m:
                trozo = trozo[m.end():]
                continue
            m = re.match(r"^(PROVINCIAS?|DISTRITOS?)( DE)?\b ?", trozo)
            if m:
                modo = "prov" if m.group(1).startswith("PROV") else "dist"
                trozo = trozo[m.end():]
                continue
            break
        if trozo and trozo not in ("Y", "DE"):
            yield modo, trozo


def desglosar(trozos, geo, preferidas):
    """Si un trozo no es un nombre conocido, prueba partirlo: 'TINTAY PUNCU YROBLE' -> 'TINTAY PUNCU' + 'ROBLE'."""
    conocido = lambda n: geo.provincia(n) or geo.distrito(n, preferidas)
    for modo, nombre in trozos:
        if conocido(nombre):
            yield modo, nombre
            continue
        w = nombre.split()
        for i in range(1, len(w)):
            izq, der = " ".join(w[:i]), " ".join(w[i:])
            der = der[1:] if der.startswith("Y") and not conocido(der) and conocido(der[1:]) else der
            if conocido(izq) and conocido(der):
                yield modo, izq
                yield modo, der
                break
        else:
            yield modo, nombre


def territorio(ruta, datos):
    """{'provincias': [...], 'distritos': [...], 'departamentos': [...], 'cobertura': {id: [dependencias]},
        'no_reconocido': [...], 'nombres': {id: nombre}}"""
    from openpyxl import load_workbook
    geo = Geo(datos)
    ws = load_workbook(ruta, read_only=True, data_only=True).worksheets[0]
    filas = ws.iter_rows(values_only=True)
    cab = [norm(c) for c in next(filas, [])]
    col = lambda *n: next((i for i, c in enumerate(cab) if any(x in c for x in n)), None)
    i_prov, i_dep, i_comp = col("PROVINCIA"), col("DEPENDENCIA"), col("COMPETENCIA")
    provs, dists, cobertura, raros = set(), set(), defaultdict(list), []
    for f in filas:
        texto = f[i_comp] if i_comp is not None and i_comp < len(f) else None
        if not texto:
            continue
        dependencia = str(f[i_dep] or "").strip() if i_dep is not None else ""
        sede = geo.provincia(norm(f[i_prov])) if i_prov is not None else None
        trozos = list(partes(texto))
        # provincias nombradas en el mismo texto o como aclaración: ayudan a ubicar distritos repetidos
        en_texto = [p for m, n in trozos if m in ("prov", "ctx") for p in [geo.provincia(n)] if p]
        trozos = [(m, n) for m, n in trozos if m != "ctx"]
        for modo, nombre in desglosar(trozos, geo, [p for p in [sede, *en_texto] if p]):
            gid = None
            if modo == "prov":
                gid = geo.provincia(nombre)
                if gid:
                    provs.add(gid)
            if not gid:
                gid = geo.distrito(nombre, [p for p in [sede, *en_texto] if p])
                if gid:
                    dists.add(gid)
            if gid:
                cobertura[gid].append(dependencia)
            else:
                raros.append(f"'{nombre.title()}' en: {dependencia}")
    # un distrito dentro de una provincia completa no hace falta listarlo aparte
    sueltos = sorted(d for d in dists if "8" + d[1:5] + "00" not in provs)
    deps = sorted({"7" + g[1:3] + "0000" for g in provs | set(sueltos)})
    return {"provincias": sorted(provs), "distritos": sueltos, "departamentos": deps,
            "cobertura": {k: sorted(set(v)) for k, v in cobertura.items()},
            "no_reconocido": sorted(set(raros)), "nombres": {k: geo.nombre.get(k, k) for k in provs | dists}}
